# -*- coding: utf-8 -*-
r"""
FilmOps Copilot — 行业标准表格数据生成 + 自动导入脚本（Demo 专用）

合规声明：
  本项目全部数据为虚构。演员、角色、经纪公司、地点、档期、通告均为合成数据，
  不包含任何真实公司、项目、艺人或个人信息。生成结果仅用于产品 Demo 演示与能力评估。

用法：
  .venv\Scripts\python scripts\make_industry_data.py            # 生成 + 注册 + 重建索引
  .venv\Scripts\python scripts\make_industry_data.py --no-index # 只生成 + 注册，不重建索引

行为（幂等，重复运行覆盖同名单）：
  1. 生成两个行业标准表格写入 data/uploaded/（内容定义见 industry_tables.py）：
     - 08_talent_schedule.xlsx 演员档期表
     - 09_daily_callsheet.xlsx 拍摄通告单
  2. 自检：预埋关键词确实写入 + 通告单到场演员 ⊆ 档期表演员（跨表一致性）
  3. 注册进 data/uploaded/manifest.json（app.ingest.register_source）
  4. 重建索引（ingest_all；--no-index 跳过）

若只想生成文件、由 UI 上传测试，请改用 make_upload_test_data.py。
"""

import argparse
import sys
from pathlib import Path

try:  # Windows 终端统一 UTF-8 输出
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from app.ingest import UPLOAD_DIR, ingest_all, parse_source, register_source  # noqa: E402
from industry_tables import (  # noqa: E402
    CALLSHEET_HEADERS, CALLSHEET_ROWS,
    TALENT_ACTORS, TALENT_HEADERS, TALENT_ROWS,
)

# ---------------------------------------------------------------------------
# 预埋内容注册表（自检用：关键词必须真实出现在生成的数据中）
# ---------------------------------------------------------------------------
PLANTED_KEYWORDS = [
    "档期待确认",          # 衔接 R8：重拍等补拍档期
    "《长风渡口》",        # 档期冲突：新剧进组 vs 补拍
    "品牌A道具未到场",     # 置入镜头顺延（呼应 R5 品牌 A 露出）
    "顺延至9-10补拍",
    "杀青宴 9-26",         # 呼应 01_schedule S-003
]

written_texts = []  # 所有写入表格的文本，供自检


def write_xlsx(filename, title, headers, rows, widths):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in rows:
        ws.append(row)
        written_texts.extend(str(v) for v in row)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    path = UPLOAD_DIR / filename
    wb.save(path)
    return path, len(rows)


def self_check():
    """预埋关键词自检 + 跨表一致性：通告单到场演员必须全部出现在档期表。"""
    all_text = "\n".join(written_texts)
    missing = [kw for kw in PLANTED_KEYWORDS if kw not in all_text]
    if missing:
        raise SystemExit(f"[FAIL] 预埋内容未写入数据: {missing}")

    call_actors = set()
    for row in CALLSHEET_ROWS:
        for name in row[6].split("、"):
            if name != "全组":
                call_actors.add(name)
    unknown = call_actors - TALENT_ACTORS
    if unknown:
        raise SystemExit(f"[FAIL] 通告单演员不在档期表中: {unknown}")
    print("[OK] 预埋内容自检通过:", ", ".join(PLANTED_KEYWORDS))
    print(f"[OK] 跨表一致性自检通过: 通告单 {len(call_actors)} 名演员全部在档期表中")


def main():
    ap = argparse.ArgumentParser(description="生成行业标准表格数据并导入索引")
    ap.add_argument("--no-index", action="store_true", help="只生成+注册，跳过重建索引")
    args = ap.parse_args()

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    results = [
        write_xlsx("08_talent_schedule.xlsx", "演员档期表", TALENT_HEADERS, TALENT_ROWS,
                   [12, 16, 8, 20, 12, 12, 12, 10, 10, 40]),
        write_xlsx("09_daily_callsheet.xlsx", "拍摄通告单", CALLSHEET_HEADERS, CALLSHEET_ROWS,
                   [12, 14, 8, 6, 10, 24, 18, 8, 12, 12, 30]),
    ]
    print(f"数据目录: {UPLOAD_DIR}")
    for path, count in results:
        print(f"  - {path.name}  ({count} 行)")

    self_check()

    # 注册进 manifest（同名覆盖，幂等）→ 解析预览
    for file_name, doc_type in [("08_talent_schedule.xlsx", "talent"),
                                ("09_daily_callsheet.xlsx", "callsheet")]:
        item = register_source(file_name, "excel", doc_type, default_dept="制片部")
        entries = parse_source(item)
        print(f"[注册] {file_name}（doc_type={doc_type}）解析出 {len(entries)} 条 Entry")

    if args.no_index:
        print("已跳过重建索引（--no-index）：数据已注册，到 ②索引管理 重建后生效")
        return

    print("重建索引中（需调用 Embedding API，请稍候）...")
    stats = ingest_all()
    print(f"[索引] Entry {stats['entries']} 条 → Chunk {stats['chunks']} 个 "
          f"（分布: {stats['chunk_types']}）")
    print("完成：新数据可检索。可试查「演员补拍档期」「品牌A道具未到场」「档期冲突」")


if __name__ == "__main__":
    main()
